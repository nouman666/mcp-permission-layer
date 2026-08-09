"""Format-only fix for Algorithms: Times New Roman, no gaps, one full step per line."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

SRC = Path(r"c:\Users\as\Desktop\bodmas\IEEE_MCP_Paper_Style_Edited.docx")

# Same algorithm meaning; each numbered step is one complete line.
ALGO1 = [
    "Algorithm 1. Permission category inference for an MCP tool call.",
    "Require: tool name n, argument map A",
    "Ensure: permission set P, high-risk flag h",
    "1: P ← ∅; h ← false",
    "2: // Layer 1: name-based heuristics",
    "3: if n matches a filesystem-read pattern then P ← P ∪ {fs.read} end if",
    "4: if n matches a filesystem-write pattern then P ← P ∪ {fs.write} end if",
    "5: if n matches an HTTP-request pattern then P ← P ∪ {net.http} end if",
    "6: if n matches a shell-execution pattern then P ← P ∪ {code.exec} end if",
    "7: if n matches an environment-read pattern then P ← P ∪ {env.read} end if",
    "8: if n matches a process-control pattern then P ← P ∪ {process} end if",
    "9: // Layer 2: argument inspection",
    "10: for each (k, v) ∈ A do",
    "11: if v denotes a filesystem path then if writeContext(n) = true then P ← P ∪ {fs.write} else P ← P ∪ {fs.read} end if end if",
    "12: if v denotes a URL then P ← P ∪ {net.http} end if",
    "13: if v contains shell metacharacters or shell commands then P ← P ∪ {code.exec} end if",
    "14: if v matches environment-variable or secret keywords then P ← P ∪ {env.read} end if",
    "15: if v matches high-risk host indicators then P ← P ∪ {net.http, code.exec}; h ← true end if",
    "16: end for",
    "17: // Layer 3: sensitive-path elevation",
    "18: for each path-valued argument v ∈ A do",
    "19: if v matches a sensitive-path pattern (e.g., ~/.ssh/**, **/.env, **/*.pem) then P ← P ∪ {env.read}; h ← true end if",
    "20: end for",
    "21: if P = ∅ then P ← {code.exec}; h ← true end if  // fail-closed over-approximation",
    "22: return P, h",
]

ALGO2 = [
    "Algorithm 2. Policy evaluation and decision aggregation.",
    "Require: tool call τ, policy π",
    "Ensure: final decision d ∈ {ALLOW, DENY, ASK} and an audit record",
    "1: (P, h) ← InferCategories(τ)  // Algorithm 1",
    "2: D ← ∅",
    "3: for each category c ∈ P do",
    "4: r ← Lookup(π, c)",
    "5: if r is undefined then r ← DefaultDeny end if",
    "6: if r.allow = false then D[c] ← DENY else if the path or domain restrictions of r are violated then D[c] ← DENY else D[c] ← r.mode end if  // ALLOW or ASK",
    "7: end for",
    "8: // Aggregate with priority DENY ≻ ASK ≻ ALLOW",
    "9: if ∃ c ∈ P : D[c] = DENY then d ← DENY else if ∃ c ∈ P : D[c] = ASK then d ← ASK else d ← ALLOW end if",
    "10: if d = ASK and no ask-handler is configured then d ← DENY end if  // fail-closed",
    "11: Append an audit record for (τ, P, D, d)",
    "12: return d",
]


def style_run(run, *, bold=False, size=9):
    run.bold = bold
    run.italic = False
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)


def tight(paragraph, *, center=False):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.line_spacing = 1.0
    pf.first_line_indent = Pt(0)
    pf.left_indent = Pt(0)


def set_caption(paragraph, text: str):
    m = re.match(r"^(Algorithm\s+\d+\.)\s*(.*)$", text.strip())
    assert m
    for r in paragraph.runs:
        r.text = ""
    tight(paragraph, center=True)
    if paragraph.runs:
        paragraph.runs[0].text = m.group(1)
        style_run(paragraph.runs[0], bold=True)
        for r in paragraph.runs[1:]:
            r.text = ""
        r2 = paragraph.add_run(" " + m.group(2))
        style_run(r2, bold=False)
    else:
        r1 = paragraph.add_run(m.group(1))
        style_run(r1, bold=True)
        r2 = paragraph.add_run(" " + m.group(2))
        style_run(r2, bold=False)


def set_line(paragraph, text: str):
    for r in paragraph.runs:
        r.text = ""
    tight(paragraph, center=False)
    # Single Times New Roman run — no Courier, no per-keyword bold
    if paragraph.runs:
        paragraph.runs[0].text = text
        style_run(paragraph.runs[0], bold=False)
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        r = paragraph.add_run(text)
        style_run(r, bold=False)


def insert_after(paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def delete_paragraph(paragraph):
    el = paragraph._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def find_algo(doc: Document, n: int) -> int:
    for i, p in enumerate(doc.paragraphs):
        if re.match(rf"^Algorithm\s+{n}\.", p.text.strip()):
            return i
    raise ValueError(n)


def algo_end(doc: Document, start: int, stop_before: int | None) -> int:
    end = start + 1
    while end < len(doc.paragraphs):
        if stop_before is not None and end >= stop_before:
            break
        t = doc.paragraphs[end].text.strip()
        if not t:
            nxt = doc.paragraphs[end + 1].text.strip() if end + 1 < len(doc.paragraphs) else ""
            if nxt.startswith("Algorithm ") or (
                nxt
                and not nxt.startswith(("Require:", "Ensure:", "Input:", "Output:"))
                and not re.match(r"^\d+:", nxt)
                and not nxt.startswith(("//", "end ", "else"))
            ):
                break
            end += 1
            continue
        if re.match(r"^Algorithm\s+\d+\.", t):
            break
        if not (
            t.startswith(("Require:", "Ensure:", "Input:", "Output:"))
            or re.match(r"^\d+:", t)
            or t.startswith(("end if", "end for", "//", "else"))
        ):
            break
        end += 1
    return end


def replace_block(doc: Document, start: int, lines: list[str], stop_before: int | None):
    end = algo_end(doc, start, stop_before)
    caption = doc.paragraphs[start]
    for i in range(end - 1, start, -1):
        delete_paragraph(doc.paragraphs[i])
    set_caption(caption, lines[0])
    prev = caption
    for line in lines[1:]:
        para = insert_after(prev)
        set_line(para, line)
        prev = para
    # one small gap after algorithm only
    gap = insert_after(prev)
    tight(gap)
    gap.paragraph_format.space_after = Pt(6)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    doc = Document(str(SRC))

    a2 = find_algo(doc, 2)
    replace_block(doc, a2, ALGO2, None)
    a1 = find_algo(doc, 1)
    a2b = find_algo(doc, 2)
    replace_block(doc, a1, ALGO1, a2b)

    doc.save(str(SRC))
    print("Saved", SRC)
    d = Document(str(SRC))
    for i, p in enumerate(d.paragraphs):
        t = p.text.strip()
        if t.startswith("Algorithm 1."):
            on = True
        if "on" in dir() and on:
            fonts = {r.font.name for r in p.runs if r.text}
            print(i, "font=", fonts, "sa=", p.paragraph_format.space_after, "|", t)
            if t.startswith("The Python") or t.startswith("IV."):
                break


if __name__ == "__main__":
    main()

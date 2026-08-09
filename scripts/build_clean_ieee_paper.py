"""
Build a clean, expanded IEEE paper (~15+ pages) from the figure-rich base.
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FIG_PNG = ROOT / "figures" / "png"
FIG_SVG = ROOT / "figures" / "svg"
# Author-finalized SVGs (do not replace with regenerated diagram code)
USER_FIG_DIR = Path(r"c:\Users\as\Desktop\bodmas\diagrasvg")
USER_FIGURE_SVGS = [
    "Fig. 1. System architecture of the runtime permission enforcement layer..svg",
    "Fig. 2. Five-stage decision pipeline..svg",
    "Fig. 3. Three-layer category inference with fail-closed default..svg",
    "diagram.svg",  # Fig. 4 decision flow
    "diagram (1).svg",  # Fig. 5 eval pipeline
    "diagram (2).svg",  # Fig. 6 ASR/TSR
    "diagram (4).svg",  # Fig. 7 trade-off
]
FIGURE_NAMES = [
    "fig1_architecture",
    "fig2_pipeline",
    "fig3_inference",
    "fig4_decision_flow",
    "fig5_eval_pipeline",
    "fig6_asr_tsr",
    "fig7_tradeoff",
]
OUTS = [
    Path(r"c:\Users\as\Desktop\bodmas\CLEAN_FINAL_IEEE_PAPER_MCP.docx"),
    Path(r"c:\Users\as\Desktop\bodmas\CLEAN_FINAL_IEEE_PAPER_MCP_MERGED.docx"),
    Path(r"c:\Users\as\Desktop\bodmas\CLEAN_FINAL_IEEE_PAPER_MCP_MERGED_NEW.docx"),
    Path(r"c:\Users\as\Desktop\bodmas\IEEE_MCP_Paper_IEEE_Final_Audited (1).docx"),
    Path(r"c:\Users\as\Desktop\bodmas\IEEE_MCP_Paper_IEEE_Final_Audited.docx"),
    Path(r"c:\Users\as\Desktop\CLEAN_FINAL_IEEE_PAPER_MCP_MERGED.docx"),
    Path(
        r"C:\Users\as\PycharmProjects\PythonResearchProject_comparision"
        r"\_mcp_zip_extract\CLEAN_FINAL_IEEE_PAPER_MCP_MERGED.docx"
    ),
    Path(r"c:\Users\as\Desktop\bodmas\diagrasvg\IEEE_MCP_Paper_IEEE_Final_Audited (1).docx"),
]


def _svg_to_png_bytes(svg_path: Path) -> bytes:
    """Rasterize an author SVG via headless Chromium (Word embeds PNG)."""
    import tempfile

    from playwright.sync_api import sync_playwright

    svg_text = svg_path.read_text(encoding="utf-8", errors="ignore")
    html = (
        "<!doctype html><html><head><meta charset='utf-8'>"
        "<style>html,body{margin:0;padding:16px;background:#fff;}"
        "svg{display:block;max-width:100%;height:auto;}</style></head><body>"
        f"{svg_text}</body></html>"
    )
    with tempfile.TemporaryDirectory() as tmp:
        html_path = Path(tmp) / "fig.html"
        html_path.write_text(html, encoding="utf-8")
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page(viewport={"width": 1400, "height": 1000})
            page.set_default_timeout(60_000)
            page.goto(html_path.resolve().as_uri(), wait_until="domcontentloaded")
            page.wait_for_timeout(200)
            loc = page.locator("svg").first
            png = loc.screenshot(type="png", timeout=60_000)
            browser.close()
    return png


def load_svg_figures() -> list[bytes]:
    """Prefer author-finalized SVGs from diagrasvg; keep captions as color keys."""
    user_paths = [USER_FIG_DIR / name for name in USER_FIGURE_SVGS]
    if all(p.exists() for p in user_paths):
        imgs: list[bytes] = []
        cache = USER_FIG_DIR / "_png_cache"
        cache.mkdir(parents=True, exist_ok=True)
        for i, svg in enumerate(user_paths, start=1):
            cached = cache / f"user_fig{i}.png"
            if not cached.exists() or cached.stat().st_mtime < svg.stat().st_mtime:
                cached.write_bytes(_svg_to_png_bytes(svg))
                print("Rasterized author SVG:", svg.name, "->", cached.name)
            imgs.append(cached.read_bytes())
        return imgs

    missing = [n for n in FIGURE_NAMES if not (FIG_PNG / f"{n}.png").exists()]
    if missing or not (FIG_SVG / f"{FIGURE_NAMES[0]}.svg").exists():
        import importlib.util
        import sys

        gen_path = Path(__file__).resolve().parent / "generate_svg_figures.py"
        spec = importlib.util.spec_from_file_location("generate_svg_figures", gen_path)
        mod = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        sys.modules["generate_svg_figures"] = mod
        spec.loader.exec_module(mod)
        mod.main()
    imgs = []
    for name in FIGURE_NAMES:
        png = FIG_PNG / f"{name}.png"
        svg = FIG_SVG / f"{name}.svg"
        if not png.exists():
            raise FileNotFoundError(f"Missing figure PNG: {png}")
        if not svg.exists():
            raise FileNotFoundError(f"Missing figure SVG: {svg}")
        imgs.append(png.read_bytes())
    return imgs


def style_run(run, size=10, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_p(
    doc,
    text,
    *,
    size=10,
    bold=False,
    italic=False,
    center=False,
    space_after=8,
    space_before=0,
    first_line=True,
    line_spacing=1.15,
):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line_spacing
    if first_line and not center and not bold:
        pf.first_line_indent = Inches(0.25)
    else:
        pf.first_line_indent = Inches(0)
    run = p.add_run(text)
    style_run(run, size=size, bold=bold, italic=italic)
    return p


def add_h(doc, text, level=1):
    size = 12 if level == 1 else 11
    return add_p(
        doc,
        text,
        size=size,
        bold=True,
        space_after=10,
        space_before=14,
        first_line=False,
    )


def add_caption(doc, text):
    return add_p(
        doc,
        text,
        size=9,
        bold=True,
        center=True,
        space_after=10,
        space_before=6,
        first_line=False,
    )


def add_fig(doc, blob: bytes | None, caption: str, width=6.0):
    if blob:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.space_before = Pt(8)
        run = p.add_run()
        run.add_picture(io.BytesIO(blob), width=Inches(width))
    add_caption(doc, caption)


def add_algo_box(doc, title: str, lines: list[str]):
    """Render numbered pseudocode in a compact, consistent IEEE-like listing."""
    add_caption(doc, title)
    for line in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.space_after = Pt(1)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.05
        pf.first_line_indent = Inches(0)
        pf.left_indent = Inches(0.15)
        run = p.add_run(line)
        style_run(run, size=9, bold=False, italic=False)
        run.font.name = "Courier New"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(10)


def add_table(doc, caption, headers, rows):
    add_caption(doc, caption)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        style_run(r, size=8, bold=True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(val))
            style_run(r, size=8)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def build():
    imgs = load_svg_figures()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.85)
    sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)

    # Title
    add_p(
        doc,
        "A Runtime Permission Enforcement Layer for Mitigating Tool Poisoning "
        "Attacks in MCP-Based Agentic Systems",
        size=16,
        bold=True,
        center=True,
        space_after=10,
        first_line=False,
    )

    # IEEE author block
    add_p(
        doc,
        "Mohammad Nouman¹*, Syed Ijlal Haider¹, and Rajja Ujjan¹",
        size=11,
        bold=False,
        center=True,
        space_after=4,
        first_line=False,
    )
    add_p(
        doc,
        "¹School of Computing, Engineering and Physical Sciences, "
        "University of the West of Scotland, Paisley, United Kingdom",
        size=9,
        italic=True,
        center=True,
        space_after=4,
        first_line=False,
    )
    add_p(
        doc,
        "*Corresponding author: Mohammad Nouman (mohammadnouman666@gmail.com)",
        size=9,
        italic=True,
        center=True,
        space_after=12,
        first_line=False,
    )

    # Abstract
    add_h(doc, "Abstract")
    add_p(
        doc,
        "Tool poisoning manipulates Model Context Protocol (MCP) tool metadata or arguments "
        "so that agentic systems perform unauthorized file access, data exfiltration, or code "
        "execution while appearing to follow legitimate user intent. Popular MCP clients largely "
        "rely on approval dialogs and static allowlists and provide little runtime least-privilege "
        "enforcement after a tool or server has been approved. We present a client-side runtime "
        "permission enforcement layer that intercepts every tool call, infers a compact set of six "
        "permission categories from tool names and arguments, and applies declarative default-deny "
        "YAML profiles with ALLOW, DENY, and ASK outcomes—without modifying existing MCP clients "
        "or servers. The layer is deployed as a stdio mediation proxy that forwards only permitted "
        "calls to upstream servers and emits append-only audit logs.",
    )
    add_p(
        doc,
        "We evaluate a multi-stage suite that separates template stress from held-out validity: an "
        "87-case pilot used only for design freeze, a 20,000-case template stress corpus (scalability "
        "only), plus remapped-name, organic, adaptive, live MCP, multi-step, and ASK studies. ASR is "
        "the fraction of malicious tool calls the layer allows to the upstream boundary (not end-to-end "
        "LLM compromise). Under limited-network: pilot ASR 1.64% (TSR 76.92%), stress ASR 2.92% "
        "(TSR 78.74%), and observed ASR 0% on organic (50/40) and live (12/13) hold-outs (small-n "
        "upper bounds in Section VI). The contribution is runtime least-privilege enforcement, not ML "
        "detection; it outperforms name-allowlist, definition-scan, and gateway-class baselines at "
        "sub-millisecond decision latency.",
    )
    add_p(
        doc,
        "Index Terms— Model Context Protocol, tool poisoning, agentic systems, runtime enforcement, "
        "least privilege, LLM security, permission models, MCP proxy.",
        italic=True,
        first_line=False,
        space_after=12,
    )

    # I. Introduction
    add_h(doc, "I. Introduction")
    add_p(
        doc,
        "Large language model (LLM) agents are increasingly entrusted with concrete host-side "
        "actions: reading and writing files, calling remote APIs, inspecting environment variables, "
        "and executing shell commands [10], [11]. The Model Context Protocol (MCP) has emerged as a "
        "practical standard for discovering and invoking such tools across heterogeneous clients and "
        "servers [12]. Standardization improves interoperability and developer velocity, but it also "
        "concentrates risk: tool names, descriptions, schemas, and runtime arguments are injected into "
        "model context and become a high-leverage control surface for attackers [1], [5].",
    )
    add_p(
        doc,
        "Tool poisoning exploits that surface. An adversary influences metadata or parameters so that "
        "the agent performs harmful side effects—credential theft from SSH or cloud directories, "
        "exfiltration over HTTP, destructive writes, or reverse-shell style execution—while the "
        "interaction still looks like ordinary tool use [2], [3]. Unlike classical malware that must "
        "bypass OS controls directly, poisoned tools co-opt the agent’s own privileged workflow. "
        "Because many clients treat an approved server as broadly trusted, a single compromised "
        "definition or argument can unlock powerful capabilities for the remainder of a session.",
    )
    add_p(
        doc,
        "Empirical evidence shows the threat is practical. MCPTox constructed 1,312 malicious cases "
        "over 45 real MCP servers and reported attack success rates up to 72.8%, with spontaneous "
        "model refusal often below 3% [2]. MCP-ITP further studied implicit poisoning, in which a "
        "poisoned tool may remain uninvoked while steering the agent toward a high-privilege "
        "legitimate call, with success up to 84.2% [3]. OWASP’s MCP Top 10 lists tool poisoning among "
        "priority risks for production deployments [5]. Collectively, these results indicate that "
        "relying on model judgment alone is insufficient.",
    )
    add_p(
        doc,
        "Despite this evidence, popular clients such as Cursor, Claude Desktop, Cline, and Continue "
        "still center their UX on approval dialogs and static allowlists [18]–[21]. After a user approves "
        "a server or tool family, fine-grained runtime least privilege at the per-call boundary is "
        "typically absent: category inference from arguments, sensitive-path elevation, and "
        "default-deny aggregation are not first-class controls. Related defenses exist—server "
        "manifests and sandboxes [4], gateway scanning and interface firewalls [8], [9], [15], and "
        "classical least-privilege principles [13], [14]—yet a lightweight, zero-modification, "
        "client-side gate that is argument-aware and reports both attack reduction and legitimate-task "
        "success remains under-explored.",
    )
    add_p(
        doc,
        "This paper addresses that gap with a systems contribution: runtime, argument-aware, "
        "client-side least-privilege enforcement for MCP tool calls. We do not propose a learned "
        "attack classifier; the design is a deterministic inference-and-policy gate. The layer "
        "(i) intercepts every tools/call request, (ii) infers permission categories using name "
        "heuristics, argument inspection, and sensitive-path elevation, (iii) evaluates declarative "
        "YAML profiles under default-deny semantics with ALLOW/DENY/ASK outcomes, and (iv) deploys "
        "as a stdio MCP proxy that requires no source changes to clients or upstream servers. "
        "Formal Algorithms 1–2 capture inference and policy aggregation exactly as implemented.",
    )
    add_p(
        doc,
        "Contributions. (1) A behavior-mapped six-category permission model with fail-closed defaults "
        "and formal Algorithms 1–2 for inference and aggregation. (2) A zero-modification runtime "
        "enforcement architecture (proxy, policy engine, audit) for MCP tool calls. (3) A multi-stage "
        "evaluation that separates template stress from hold-out validity (organic, remapped-name, "
        "adaptive, live MCP, ASK), reporting ASR, TSR, Attack-Blocking F1, latency, and bootstrap "
        "CIs. (4) A reproducible artifact package for independent verification.",
    )

    # II. Literature Review
    add_h(doc, "II. Literature Review")
    add_h(doc, "A. Indirect Prompt Injection and Tool-Oriented Attacks", level=2)
    add_p(
        doc,
        "Greshake et al. formalized indirect prompt injection against LLM-integrated applications, "
        "showing that malicious instructions embedded in retrieved or otherwise untrusted content can "
        "compromise systems without a direct attacker chat interface [1]. Subsequent agent-security "
        "benchmarks such as InjecAgent and AgentDojo demonstrated that tool-using agents remain "
        "vulnerable when untrusted observations influence tool selection and arguments [6], [7]. "
        "These results established that the boundary between “data” and “instructions” is porous once "
        "models can act through tools.",
    )
    add_p(
        doc,
        "Tool poisoning specializes this threat to MCP. Explicit-trigger hijacking places overt harmful "
        "instructions in tool metadata. Implicit-trigger hijacking biases the agent toward unsafe tool "
        "sequences without an overt command. Parameter tampering alters paths, URLs, recipients, or "
        "command strings so that a seemingly legitimate tool produces harmful effects [2], [3]. Because "
        "MCP servers advertise capabilities that clients then trust, poisoning can occur at registration "
        "time, at schema presentation time, or at call time through attacker-controlled arguments.",
    )
    add_h(doc, "B. MCP Threat Benchmarks", level=2)
    add_p(
        doc,
        "MCPTox provides a large attack corpus grounded in real MCP servers and reports high ASR under "
        "multiple models [2]. MCP-ITP automates implicit poisoning scenarios and shows that defenders "
        "cannot assume the poisoned tool is the one that executes [3]. Broader MCP landscape and "
        "threat surveys similarly warn that tool ecosystems accumulate untrusted capability "
        "descriptions over time [22]. These benchmarks primarily measure attack success; "
        "they generally do not evaluate a client-side runtime permission gate with paired utility "
        "metrics, which is the evaluation niche this paper targets.",
    )
    add_h(doc, "C. Defenses: Server, Gateway, and Client Controls", level=2)
    add_p(
        doc,
        "Server-oriented defenses such as AgentBound attach Android-style manifests and sandboxes to "
        "MCP servers, constraining what a server process may do on the host [4], [13]. Gateway and "
        "control-plane systems scan definitions, pin schemas, apply allowlists, or minimize tool "
        "interfaces exposed to the model [8], [9], [15]. Interface firewalls that aggressively "
        "reduce tool inputs can be highly effective on some injection benchmarks [9]. Client products "
        "still emphasize human approval and static name allowlists [18]–[21]. Complementary lines of "
        "work study design-time hardening of agent architectures against prompt injection [16] and "
        "broader LLM application risks catalogued by OWASP [17].",
    )
    add_p(
        doc,
        "Classical least privilege motivates mapping actions to named categories, requiring explicit "
        "grants, and denying by default [13], [14]. In the agent setting, permission models must also "
        "cope with natural-language tool names, free-form arguments, and rapidly changing server "
        "inventories. Server sandboxing does not replace a client-side gate over categories and "
        "arguments; definition scanning alone does not enforce least privilege on every later call; "
        "and model refusal remains unreliable under metadata manipulation [2], [3].",
    )
    add_h(doc, "D. Gap and Positioning", level=2)
    add_p(
        doc,
        "Table I compares representative prior work along dimensions that matter for practical MCP "
        "deployment: client-side placement, runtime category-and-argument checks, zero modification of "
        "existing software, and reporting of legitimate-task utility alongside security. The remaining "
        "gap is a lightweight client-side, argument-aware runtime gate that needs no MCP source changes "
        "and reports both attack reduction and task success. This paper occupies that cell.",
    )
    add_table(
        doc,
        "TABLE I. LITERATURE REVIEW COMPARISON",
        ["Work", "Focus", "Client-side", "Runtime cat.+args", "Zero modify", "Reports utility"],
        [
            ["Greshake et al. [1]", "Indirect prompt injection", "N/A", "No", "N/A", "No"],
            ["MCPTox [2]", "MCP tool-poisoning bench", "N/A", "No (attack)", "N/A", "No"],
            ["MCP-ITP [3]", "Implicit tool poisoning", "N/A", "No (attack)", "N/A", "No"],
            ["AgentBound [4]", "Server manifests/sandbox", "No", "Server perms", "Server-side", "Limited"],
            ["Gateway/firewall [8][9][15]", "Scan/allowlist/minimize", "Partial", "Partial", "Varies", "Often no"],
            ["Client allowlists [18]–[21]", "Approval + static names", "Yes", "No (static)", "N/A", "Rare"],
            ["This work", "Runtime permission layer", "Yes", "Yes", "Yes", "Yes"],
        ],
    )

    # III. Design
    add_h(doc, "III. System Design")
    add_h(doc, "A. Threat Model and Assumptions", level=2)
    add_p(
        doc,
        "We assume an adversary who can supply or influence MCP tool metadata (names, descriptions, "
        "schemas) and/or call arguments, including by registering or pointing the client at a malicious "
        "MCP server. The adversary may craft multi-parameter payloads, rename tools to evade naive "
        "string matches, and, in the adaptive setting, know the six categories, path heuristics, and "
        "domain rules used by the defender. The adversary’s goal is harmful host-side effects: reading "
        "secrets, writing attacker-controlled content, contacting attacker domains, or executing commands.",
    )
    add_p(
        doc,
        "The adversary cannot modify the enforcement proxy binary or source, rewrite policy files, "
        "tamper with the append-only audit logger, compromise the client OS trust boundary, or silently "
        "disable configuration that routes tools/call traffic through the proxy. We assume the proxy, "
        "policies, and logger execute in a trusted client-side boundary in the sense of classical "
        "protection of information systems [14]. Social engineering that induces a user to weaken "
        "profiles (for example, allow_domains: [\"*\"]) is treated as misconfiguration risk rather than "
        "an in-scope bypass of a correctly deployed gate.",
    )
    add_h(doc, "B. Permission Taxonomy Derivation", level=2)
    add_p(
        doc,
        "We derived categories with a structured, evidence-driven procedure rather than an a priori "
        "ontology. Step 1: collect harmful outcomes reported in MCPTox, MCP-ITP, and OWASP MCP Top 10 "
        "[2], [3], [5] (credential theft, exfiltration, shell execution, env/secret access, destructive "
        "writes, process control). Step 2: map each outcome to the minimal host capability required to "
        "realize it. Step 3: merge capabilities that share the same enforcement predicate (for example, "
        "SSH/cloud file reads both need filesystem read plus sensitive-path elevation). Step 4: stop "
        "when adding a category would not change ALLOW/DENY decisions on the pilot threat set under "
        "default-deny profiles. Table II records the resulting behavior→capability→category map. The "
        "six categories—fs.read, fs.write, net.http, code.exec, env.read, and process—therefore form "
        "a compact set covering the dominant tool-poisoning behaviors evaluated in this study; they "
        "are not claimed to cover all MCP capabilities (database APIs, browser automation, cloud IAM, "
        "email, arbitrary SaaS tools). Overlaps are handled by set union: a call may require multiple "
        "categories, and Algorithm 2 aggregates with Deny ≻ Ask ≻ Allow.",
    )
    add_p(
        doc,
        "Two design choices are intentional. First, sensitive filesystem reads (for example, ~/.ssh/** "
        "or **/.env) elevate to env.read and mark the call high-risk, because credential-shaped reads "
        "are closer to secret access than to ordinary document reading. Second, if no category is "
        "inferred, the prototype currently fails closed by assigning code.exec with high_risk. This "
        "is a security over-approximation, not a claim that the unknown tool is literally a shell: "
        "for example, weather_lookup(city=\"London\") would be treated as high-risk under that rule. "
        "Ablation confirms the utility cost (TSR falls to 42.31% without name heuristics). A cleaner "
        "future design is a first-class unknown category with an explicit policy rule, separating "
        "semantic classification from fail-closed policy; we discuss this in Section VII-D.",
    )
    add_table(
        doc,
        "TABLE II. TAXONOMY DERIVATION FROM ATTACK BEHAVIORS",
        ["Attack behavior", "Required capability", "Category"],
        [
            ["Read SSH/cloud secrets", "Filesystem read of sensitive path", "fs.read + elevation"],
            ["Exfiltrate file contents", "Read + outbound HTTP", "fs.read + net.http"],
            ["Execute shell / reverse shell", "Command execution", "code.exec"],
            ["Read environment secrets", "Environment access", "env.read"],
            ["Overwrite configs / keys", "Filesystem write", "fs.write"],
            ["Spawn/kill processes", "Process control", "process"],
        ],
    )
    add_h(doc, "C. Policy Profiles", level=2)
    add_p(
        doc,
        "Profiles are declarative YAML documents that specify, per category, whether the action is "
        "allowed, denied, or requires ASK, plus optional path and domain restrictions. Four profiles "
        "span the security–utility spectrum used throughout evaluation (Table III). Read-only permits "
        "ordinary filesystem reads while denying writes, execution, and unconstrained networking. "
        "Limited-network additionally permits HTTP to an explicit domain allow-list and is our primary "
        "recommended default for interactive coding agents. No-code-exec preserves broader filesystem "
        "and network utility while denying execution. Approval-gated routes high-risk categories to ASK "
        "so a human can restore utility without permanently widening the policy.",
    )
    add_p(
        doc,
        "Open network allow-lists (for example, allow_domains: [\"*\"]) can nullify net.http "
        "restrictions and should be treated as high-risk misconfiguration. Profiles are hot-loadable "
        "so operators can tighten policy after an incident without redeploying clients or servers.",
    )
    add_table(
        doc,
        "TABLE III. POLICY PROFILES",
        ["Profile", "Allowed", "Main restriction"],
        [
            ["read-only", "fs.read", "Sensitive paths denied"],
            ["limited-network", "fs.read, net.http", "Domain allow-list"],
            ["no-code-exec", "fs.read/write, net.http", "code.exec denied"],
            ["approval-gated", "Most; high-risk → ASK", "Human confirm on risk"],
        ],
    )
    add_h(doc, "D. Architecture and Inference", level=2)
    add_p(
        doc,
        "The system comprises four components. The Policy Manager loads and validates YAML profiles. "
        "The Permission Checker runs category inference and policy evaluation. The Enforcement Proxy "
        "exposes both an in-process Python API and a stdio MCP mediator that sits between the client "
        "and upstream servers. The Audit Logger appends structured JSON Lines records for each "
        "decision. Logs may contain tool names and argument summaries and should be access-controlled; "
        "logging cost is negligible relative to model and network latency.",
    )
    add_p(
        doc,
        "Fig. 1 shows the placement of the layer in a typical MCP deployment. Fig. 2 summarizes the "
        "five-stage pipeline: receive call, infer categories, evaluate policy, aggregate decision, and "
        "enforce/log. Fig. 3 highlights the three-layer inference design with fail-closed default. "
        "In Figs. 1–5, fill colors are role-coded as stated in each figure caption (client/input, "
        "enforcement, checking/ALLOW, inference/ASK, policy, upstream/decision, audit/DENY, and "
        "persistent stores). In Figs. 6–7, colors identify policy profiles, not roles. "
        "Because mediation occurs on the client host before upstream execution, denied calls never "
        "reach potentially malicious or over-privileged servers.",
    )
    add_fig(
        doc,
        imgs[0],
        "Fig. 1. System architecture of the runtime permission enforcement layer. "
        "Color key: blue=client, orange=proxy, green=checker, purple=server, teal=policy manager, "
        "yellow=inference, salmon=audit, grey=persistent store; red arrows=audit path.",
    )
    add_fig(
        doc,
        imgs[1],
        "Fig. 2. Five-stage decision pipeline. "
        "Colors mark successive stages (parse → infer → policy → restriction → aggregate).",
    )
    add_fig(
        doc,
        imgs[2],
        "Fig. 3. Three-layer category inference with fail-closed default. "
        "Color key: blue=input call, yellow=Layer 1 name heuristics, orange=Layer 2 argument "
        "inspection, salmon=Layer 3 sensitive elevation.",
    )

    add_h(doc, "E. Formal Algorithms", level=2)
    add_p(
        doc,
        "Algorithm 1 returns the permission set required by a tool call. Layer 1 applies name "
        "heuristics; Layer 2 inspects argument values for paths, URLs, shell metacharacters, secret "
        "keywords, and attackerish host markers; Layer 3 elevates sensitive paths. If the resulting "
        "set is empty, the algorithm fails closed to code.exec. Algorithm 2 maps that set to a final "
        "ALLOW / DENY / ASK decision under default-deny semantics, with Deny outranking Ask, which "
        "outranks Allow. If ASK is selected but no ask_handler is configured, the implementation "
        "fails closed to DENY. Both algorithms are implemented verbatim in the prototype and exercised "
        "by unit tests and the evaluation harness.",
    )
    add_p(
        doc,
        "Exact matching rules ship in src/inference.py of the artifact package. Name heuristics use "
        "curated token lists (for example, read_file/list_dir/glob → fs.read; write_file/delete_file → "
        "fs.write; run_terminal/bash/exec → code.exec; http_request/fetch/download → net.http). A value "
        "looks like a path if it starts with /, ~, ./, ../, or a Windows drive prefix, or contains a "
        "slash plus a known extension. Sensitive-path elevation matches patterns such as ~/.ssh/**, "
        "~/.aws/**, **/.env, **/*.pem, and credential/secret filename markers. Attackerish markers "
        "include substrings such as attacker, evil.com, and exfil. Secret-keyword checks include "
        "api_key, token, password, and private_key. These lists are intentionally over-approximate and "
        "versioned with the code rather than restated as an exhaustive appendix here.",
    )
    add_algo_box(
        doc,
        "Algorithm 1 Permission Inference",
        [
            "Input: tool_name, arguments (map)",
            "Output: categories P, high_risk flag",
            "1:  P ← ∅;  high_risk ← false",
            "2:  // Layer 1 — name heuristics",
            "3:  if name matches read/list/glob patterns then P ← P ∪ {fs.read}",
            "4:  if name matches write/delete/mkdir patterns then P ← P ∪ {fs.write}",
            "5:  if name matches http/fetch/request/email patterns then P ← P ∪ {net.http}",
            "6:  if name matches terminal/bash/exec/shell patterns then P ← P ∪ {code.exec}",
            "7:  if name matches get_env/environ patterns then P ← P ∪ {env.read}",
            "8:  if name matches process/kill/spawn patterns then P ← P ∪ {process}",
            "9:  // Layer 2 — argument inspection",
            "10: for all (key, value) in arguments do",
            "11:   if value looks like a path then",
            "12:     if write-context(tool_name) then P ← P ∪ {fs.write}",
            "13:     else P ← P ∪ {fs.read}",
            "14:   if value looks like a URL then P ← P ∪ {net.http}",
            "15:   if value has shell metacharacters / shell commands then P ← P ∪ {code.exec}",
            "16:   if value matches env-var / secret keywords then P ← P ∪ {env.read}",
            "17:   if value matches attackerish host markers then",
            "18:     P ← P ∪ {net.http, code.exec};  high_risk ← true",
            "19: // Layer 3 — sensitive-path elevation",
            "20: for all path-like values v in arguments do",
            "21:   if v matches ~/.ssh/**, ~/.aws/**, **/.env, **/*.pem, credentials, … then",
            "22:     P ← P ∪ {env.read};  high_risk ← true",
            "23: // Fail-closed unknown: over-approx. as code.exec (not a semantic claim)",
            "24: if P = ∅ then P ← {code.exec};  high_risk ← true",
            "25: return P, high_risk",
        ],
    )
    add_algo_box(
        doc,
        "Algorithm 2 Policy Evaluation and Aggregation",
        [
            "Input: tool_call, policy",
            "Output: final ∈ {ALLOW, DENY, ASK}, audit record",
            "1:  (P, high_risk) ← InferCategories(tool_call)   // Algorithm 1",
            "2:  D ← empty map",
            "3:  for all category c in P do",
            "4:    rule ← Lookup(policy, c)",
            "5:    if rule missing then rule ← default-deny",
            "6:    if rule.allow = false then D[c] ← DENY",
            "7:    else if path/domain restrictions violated then D[c] ← DENY",
            "8:    else D[c] ← rule.mode   // ALLOW or ASK",
            "9:  // Aggregate with priority Deny > Ask > Allow",
            "10: if ∃ c : D[c] = DENY then final ← DENY",
            "11: else if ∃ c : D[c] = ASK then final ← ASK",
            "12: else final ← ALLOW",
            "13: if final = ASK and no ask_handler then final ← DENY   // fail-closed",
            "14: Log(DecisionRecord(tool_call, P, D, final, latency))",
            "15: return final",
        ],
    )

    # IV. Implementation
    add_h(doc, "IV. Implementation")
    add_h(doc, "A. Prototype Organization", level=2)
    add_p(
        doc,
        "The Python 3 prototype (mcp_permission_layer) implements Algorithms 1–2 as pure functions "
        "plus policy loading, audit logging, and the proxy. Core modules cover category inference, "
        "policy lookup, decision aggregation, YAML profile parsing, and JSON Lines audit records. "
        "Ablation flags allow evaluation scripts to disable name heuristics, argument inspection, "
        "sensitive-path elevation, or domain allow-listing without rewriting policies. Thirty-three "
        "unit tests cover inference edge cases, policy defaults, checker aggregation, audit fields, "
        "and proxy allow/deny forwarding behavior.",
    )
    add_h(doc, "B. Deployment Model", level=2)
    add_p(
        doc,
        "Deployment is configuration-only. The MCP client is pointed at the enforcement proxy instead "
        "of the upstream server; the proxy advertises tools, accepts tools/call requests, runs "
        "Algorithm 2, and forwards only ALLOWED requests. DENY responses are returned locally with a "
        "structured error, preventing side effects on the host and on remote services. ASK invokes an "
        "optional handler; if none is configured, ASK fails closed to DENY. Fig. 4 shows the "
        "implemented decision flow inside the Enforcement Proxy.",
    )
    add_p(
        doc,
        "Because the proxy speaks standard MCP stdio framing, the same binary can mediate heterogeneous "
        "upstream servers (filesystem, fetch, git, shell, and custom tools) without server-specific "
        "patches. Operators select a profile per workspace or trust tier. This zero-modification "
        "property is essential for adoption in environments where clients and third-party MCP servers "
        "cannot be recompiled or forked.",
    )
    add_fig(
        doc,
        imgs[3],
        "Fig. 4. Tool-call decision flow in the Enforcement Proxy. "
        "Outcome colors: salmon=DENY, yellow=ASK, green=ALLOW, grey=audit sink; "
        "arrow colors match the chosen path (red/amber/green).",
    )

    # V. Setup
    add_h(doc, "V. Experimental Setup")
    add_h(doc, "A. Datasets", level=2)
    add_p(
        doc,
        "Evaluation uses multiple datasets with an explicit independence protocol. Phase A "
        "(design freeze): taxonomy, NAME_RULES / path patterns / attackerish markers in "
        "src/inference.py, and YAML profiles were fixed using only the hand-curated pilot_87 "
        "(61 malicious / 26 benign) plus published threat descriptions [2], [3], [5]—not the "
        "20k corpus and not the hold-outs. Phase B (stress generation): the template stress "
        "corpus (datasets/large_dataset.jsonl) was then generated by deterministic expansion of "
        "seed templates (sensitive paths, attacker URLs, shell payloads, safe project paths) into "
        "exactly 20,000 cases (10k/10k) with 23 tool names and category counts "
        "credential_reading 3556, data_exfiltration 2160, code_execution 2160, parameter_tampering "
        "1152, cross_tool_escalation 540, logging_surveillance 432; ids are unique and expansion "
        "needs no random seed. Phase C (held-out evaluation): remapped-name, organic, adaptive, "
        "live MCP, multi-step, and ASK sets were evaluated without further heuristic edits. We "
        "therefore treat the 20k corpus strictly as a scalability/stress test of repeated patterns, "
        "not as evidence of organic diversity or as a tuning set. External-validity claims rest on "
        "Phase-C hold-outs.",
    )
    add_p(
        doc,
        "Phase-C sets: remapped-name (pilot_87 names replaced with unseen identifiers); organic "
        "hold-out (50 malicious / 40 benign, hand-authored, non-template); adaptive set (15 "
        "malicious / 5 benign: benign-looking names, args-only malice, encoding); 13 malicious "
        "multi-step chains; 25 live MCP workflows (12 malicious / 13 benign); 20 ASK prompts. "
        "All dataset files ship under datasets/.",
    )
    add_h(doc, "B. Metrics and Protocol", level=2)
    add_p(
        doc,
        "We report enforcement-boundary metrics, not full agent-compromise metrics. Attack Success "
        "Rate (ASR) is the proportion of malicious tool-call cases that the enforcement layer allows "
        "through to the upstream execution boundary (Decision=ALLOW). Block Rate = 1 − ASR on the "
        "malicious subset. Task Success Rate (TSR) is the proportion of benign cases allowed; we use "
        "TSR uniformly and do not introduce a separate “benign success” synonym. Attack-Blocking F1 "
        "(reported as F1) treats a correctly blocked malicious tool call as the positive class: "
        "TP = malicious AND DENY, FP = benign AND DENY, FN = malicious AND ALLOW, TN = benign AND "
        "ALLOW. Equivalently, F1 scores the detector that flags attacks by denying them; it is not "
        "an F1 over “allow = positive.” Unless stated otherwise, TSR is computed on the benign "
        "subset of the same dataset as ASR (for pilot_87: 26 benign cases). Latency is "
        "reported as mean/median/p95/p99 on the decision path; live MCP additionally reports "
        "end-to-end latency including upstream execution. Bootstrap 95% CIs (n=1000) are reported for "
        "limited-network / pilot_87. Automated runs map ASK→DENY unless a study enables an ASK "
        "handler. Baselines: B0 none, B1 static allowlist, B2 definition scan, B3 gateway filter, "
        "B4 full layer. Fig. 5 summarizes the pipeline.",
    )
    add_p(
        doc,
        "All automated experiments use the same harness entry points so that policy profiles, ablation "
        "flags, and dataset loaders remain comparable. Live MCP experiments exercise the real stdio "
        "proxy path rather than a mocked checker, ensuring that serialization, forwarding, and denial "
        "short-circuiting are included in end-to-end measurements.",
    )
    add_fig(
        doc,
        imgs[4],
        "Fig. 5. Evaluation pipeline used across all studies. "
        "Color key: blue=dataset, orange=policy under test, green=proxy, purple=decision, "
        "salmon=metrics.",
    )

    # VI. Results
    add_h(doc, "VI. Results")
    add_h(doc, "A. Pilot and Template Stress", level=2)
    add_p(
        doc,
        "Without enforcement, ASR is 100% by construction on malicious cases that reach the upstream "
        "boundary unprotected. On pilot_87 (61 malicious / 26 benign), limited-network yields the "
        "best measured balance (Table IV): ASR 1.64% (1/61 allowed), Block 98.36% (60/61), TSR 76.92% "
        "(20/26) on the same benign subset, and Attack-Blocking F1 94.49%. Independently: TP=60 "
        "(malicious blocked), FP=6 (benign blocked), TN=20 (benign allowed), FN=1 (malicious "
        "allowed), so Precision=60/66 and Recall=60/61, hence F1=2PR/(P+R)=94.49%, matching "
        "Table IV. The six blocked benign cases are all "
        "writes—an expected cost of a profile that does not broadly allow fs.write. Read-only is "
        "stricter on TSR (57.69%); no-code-exec preserves TSR 100% but raises ASR to 6.56%; "
        "approval-gated (ASK→DENY) sits between them.",
    )
    add_p(
        doc,
        "On the 20,000-case stress corpus (10k/10k; Table V; Fig. 6–7), limited-network reaches ASR "
        "2.92% and TSR 78.74% at about 0.52 ms mean decision-path latency. Read-only matches the same "
        "ASR but loses substantial utility. No-code-exec and approval-gated preserve more TSR but "
        "admit higher ASR (17.68% and 10.48%). Fig. 7 shows that limited-network is non-dominated in "
        "the measured security–utility trade-off. Consistent with Section V-A, we interpret these "
        "figures as stress/scalability evidence only; organic and live validity are assessed in "
        "Sections VI-C and VI-D.",
    )
    add_table(
        doc,
        "TABLE IV. PILOT SECURITY RESULTS (61 MALICIOUS / 26 BENIGN; F1 = ATTACK-BLOCKING F1)",
        ["Policy", "ASR", "Block", "TSR", "F1"],
        [
            ["baseline", "100%", "0%", "100%", "0%"],
            ["read_only", "1.64%", "98.36%", "57.69%", "90.91%"],
            ["limited_network", "1.64%", "98.36%", "76.92%", "94.49%"],
            ["no_code_exec", "6.56%", "93.44%", "100%", "96.61%"],
            ["approval_gated", "4.92%", "95.08%", "76.92%", "92.80%"],
        ],
    )
    add_table(
        doc,
        "TABLE V. TEMPLATE STRESS RESULTS (10k MAL / 10k BENIGN)",
        ["Policy", "ASR", "TSR", "F1", "ms/call"],
        [
            ["read_only", "2.92%", "57.62%", "81.08%", "0.51"],
            ["limited_network", "2.92%", "78.74%", "88.93%", "0.52"],
            ["no_code_exec", "17.68%", "97.32%", "88.99%", "0.50"],
            ["approval_gated", "10.48%", "81.38%", "86.02%", "0.74"],
        ],
    )
    add_fig(
        doc,
        imgs[5],
        "Fig. 6. ASR and TSR on the 20,000-case template stress corpus. "
        "Bar colors identify policy profiles (same color in both panels).",
    )
    add_fig(
        doc,
        imgs[6],
        "Fig. 7. Security–utility trade-off; limited-network is non-dominated. "
        "Point colors identify policy profiles; dashed lines are TSR/Block reference guides.",
    )

    add_h(doc, "B. Baselines", level=2)
    add_p(
        doc,
        "Baselines B0–B3 are not strawmen; they instantiate the defense classes that dominate "
        "today’s MCP client and gateway practice [8], [15], [18]–[21], implemented in "
        "baselines/defenses.py and scored with the same pilot_87 loader and metrics as B4 "
        "(Table VI). B0 (always ALLOW) is the unprotected reference. B1 (static allowlist) models "
        "name-based client allowlisting: it allows SAFE_TOOL_ALLOWLIST names "
        "(read_file/list_dir/glob/http_request/fetch/write_file/…) and denies others, ignoring "
        "arguments/descriptions—the common “approve these tools” UX. B2 (definition scan) models "
        "metadata scanners that reject suspicious tool descriptions/names via curated keywords "
        "(ignore previous, exfiltrat, attacker, hidden instruction, private key, | bash, …) without "
        "runtime argument checks—the class of definition-time MCP scanners. B3 (gateway filter) "
        "models lightweight call-time gateways: high-risk name fragments plus regexes on "
        "attackerish URLs, shell metacharacters, and crude sensitive-path substrings, but without a "
        "category model, YAML profiles, domain allow-lists, or Deny≻Ask≻Allow aggregation. B4 is "
        "the full limited-network EnforcementProxy (ASK→DENY in automated runs). All baselines "
        "receive identical (tool_name, arguments, description) fields. We report B1–B3 as "
        "representative class baselines, not as the strongest possible instance of each class; "
        "stronger scanners/gateways exist, but they still typically lack the combination of "
        "category inference, sensitive-path elevation, and declarative default-deny profiles that "
        "B4 evaluates.",
    )
    add_p(
        doc,
        "On pilot_87, B1 preserves TSR 100% but leaves ASR 54.10% because allowlisted names can still "
        "carry poisoned arguments. B2 reaches ASR 65.57% because many attacks place malice in "
        "parameters rather than descriptions. B3 reduces ASR to 24.59% via name/argument heuristics "
        "but still lags B4 (ASR 1.64%, TSR 76.92%, F1 94.49%). The B3–B4 gap isolates the value of "
        "structured categories plus profiled path/domain restrictions beyond a flat gateway regex.",
    )
    add_table(
        doc,
        "TABLE VI. BASELINE COMPARISON (PILOT_87; 61 MAL / 26 BENIGN; F1 = ATTACK-BLOCKING F1)",
        ["Baseline", "ASR", "Block", "TSR", "F1"],
        [
            ["B0 no defense", "100.00%", "0.00%", "100.00%", "0.00%"],
            ["B1 static allowlist", "54.10%", "45.90%", "100.00%", "62.92%"],
            ["B2 definition scan", "65.57%", "34.43%", "100.00%", "51.22%"],
            ["B3 gateway filter", "24.59%", "75.41%", "100.00%", "85.98%"],
            ["B4 full layer", "1.64%", "98.36%", "76.92%", "94.49%"],
        ],
    )

    add_h(doc, "C. Hold-out and Organic Validity", level=2)
    add_p(
        doc,
        "On the remapped-name split of pilot_87 (unseen tool identifiers; n=87), limited-network "
        "produced an observed ASR of 3.28%, compared with 2.92% on the template stress corpus "
        "(n=20,000), while TSR decreased to 65.38%. Because these evaluations use different "
        "datasets and sample sizes, we report the figures side-by-side rather than claiming a "
        "causal degradation attributable solely to renaming. The lower TSR is consistent with "
        "Layer-1 name heuristics mattering for utility: when names are unknown, fail-closed "
        "defaults over-deny some benign tools. The organic hold-out (50 malicious / 40 benign)—hand-authored, "
        "non-template, and never used for tuning—shows observed ASR 0% (0/50) and TSR 100% (40/40) "
        "under limited-network (Table VII). With zero observed successes, a one-sided 95% upper bound "
        "on the true ASR is ≤5.8% (1−0.05^(1/n)); we therefore treat 0% as an observed point estimate, "
        "not a proof of zero risk. Argument inspection, sensitive-path elevation, and fail-closed "
        "unknowns still carry much of the external-validity argument beyond template stress.",
    )
    add_table(
        doc,
        "TABLE VII. ORGANIC-STYLE HOLD-OUT (50 MAL / 40 BENIGN)",
        ["Profile", "ASR", "TSR", "F1"],
        [
            ["read_only", "0.00%", "75.00%", "90.91%"],
            ["limited_network", "0.00%", "100.00%", "100.00%"],
            ["no_code_exec", "2.00%", "100.00%", "98.99%"],
            ["approval_gated", "2.00%", "100.00%", "98.99%"],
        ],
    )

    add_h(doc, "D. Live MCP Deployment", level=2)
    add_p(
        doc,
        "To validate mediation outside synthetic harnesses, 25 workflows (12 malicious / 13 benign) "
        "were executed on the live path client → stdio proxy → real upstream servers (filesystem, "
        "fetch, git, and shell). Limited-network yields observed ASR 0% (0/12) and TSR 100% (13/13) "
        "(Table VIII); the corresponding one-sided 95% upper bound on true ASR is ≤22.1%, so this "
        "result is encouraging but not statistically conclusive alone. Decision-path latency remains "
        "sub-millisecond; previously measured end-to-end means around tens of milliseconds are "
        "upstream-dominated. Read-only keeps observed ASR 0% but reduces TSR by denying legitimate "
        "writes; no-code-exec admits ASR 8.33% (1/12) while preserving TSR 100%.",
    )
    add_table(
        doc,
        "TABLE VIII. LIVE REAL MCP DEPLOYMENT (12 MAL / 13 BENIGN; DEC. MS = DECISION-PATH ONLY)",
        ["Defense", "ASR", "TSR", "F1", "Dec. ms"],
        [
            ["B0 no defense", "100%", "100%", "0%", "n/a"],
            ["B4 limited_network", "0%", "100%", "100%", "0.45"],
            ["B4 read_only", "0%", "76.9%", "88.9%", "0.52"],
            ["B4 no_code_exec", "8.3%", "100%", "95.7%", "0.41"],
            ["B4 approval_gated", "0%", "100%", "100%", "0.39"],
        ],
    )
    add_p(
        doc,
        "Notes: (1) Under Attack-Blocking F1 (Section V: positive = correctly blocked malicious call), "
        "B0 has TP=0 and FN=12 because it never denies a malicious call, so Precision is zero and "
        "F1=0% despite TSR=100%—expected for an always-allow baseline. (2) “Dec. ms” is mean "
        "decision-path latency of the enforcement check only (not end-to-end MCP/upstream runtime); "
        "end-to-end latency remains upstream-dominated (tens of ms in prior live measurements).",
        italic=True,
        first_line=False,
        size=9,
        space_after=8,
    )

    add_h(doc, "E. Adaptive and Multi-Step Attacks", level=2)
    add_p(
        doc,
        "We treat the adaptive suite as a qualitative adversarial robustness probe, not a "
        "statistically powered benchmark: 15 malicious / 5 benign cases crafted with knowledge of "
        "Algorithm 1 (category names, path heuristics, domain rules), spanning benign-looking names, "
        "args-only malice, and light encoding. All four profiles show observed ASR 0% (0/15; "
        "one-sided 95% upper bound ≤18.1%; Table IX; B0 = 100%); limited-network also keeps TSR "
        "100% (5/5). These results motivate larger adaptive corpora in future work.",
    )
    add_p(
        doc,
        "Multi-step evaluation uses 13 malicious chains. Formal success definition: a chain fully "
        "succeeds iff every step is ALLOW under the evaluated mediator—equivalently, "
        "success_requires=all_allowed in the dataset—so secret acquisition, transfer/sink execution, "
        "and any attacker-controlled effect encoded in the chain all complete without a DENY. "
        "Partial progress (for example, 2 of 3 stages ALLOW) counts as failure for full-chain ASR; "
        "we do not claim that intermediate allows are impossible. Production enforcement is per-call. "
        "Separately, the harness includes an experimental session-taint mediator that marks "
        "high-risk/secret-read attempts and blocks later sinks (net.http, fs.write, code.exec, "
        "process). Under both mediators, 0/13 chains fully succeed. The session-taint mediator is "
        "not a production information-flow control plane (Section VII-D).",
    )
    add_table(
        doc,
        "TABLE IX. ADAPTIVE PROBE (15 MAL / 5 BENIGN; QUALITATIVE)",
        ["Profile", "ASR", "TSR", "F1"],
        [
            ["read_only", "0.00%", "60.00%", "93.75%"],
            ["limited_network", "0.00%", "100.00%", "100.00%"],
            ["no_code_exec", "0.00%", "100.00%", "100.00%"],
            ["approval_gated", "0.00%", "100.00%", "100.00%"],
            ["B0 no defense", "100.00%", "100.00%", "0.00%"],
        ],
    )

    add_h(doc, "F. Ablation", level=2)
    add_p(
        doc,
        "Component knock-outs on limited-network / pilot_87 quantify which mechanisms matter "
        "(Table X). Removing argument inspection raises ASR from 1.64% to 4.92%, confirming that "
        "parameter-borne malice is not fully covered by names alone; the same mechanism is what keeps "
        "remapped-name ASR low (3.28%) when Layer-1 tokens disappear. Removing sensitive-path "
        "elevation raises ASR to 14.75%, showing that ordinary fs.read permission is insufficient for "
        "credential-shaped paths. Removing domain allow-listing raises ASR to 4.92%. Removing name "
        "heuristics collapses TSR to 42.31% because fail-closed unknowns over-deny benign tools whose "
        "arguments alone do not yet reveal intent. Thus argument inspection is a security contributor "
        "on pilot and remapped settings, while name heuristics are primarily a utility contributor.",
    )
    add_table(
        doc,
        "TABLE X. ABLATION (LIMITED-NETWORK, PILOT_87)",
        ["Config", "ASR", "TSR", "F1"],
        [
            ["Full system", "1.64%", "76.92%", "94.49%"],
            ["− argument inspection", "4.92%", "76.92%", "92.80%"],
            ["− sensitive-path elevation", "14.75%", "76.92%", "87.39%"],
            ["− domain allow-list", "4.92%", "76.92%", "92.80%"],
            ["− name heuristics", "1.64%", "42.31%", "88.24%"],
        ],
    )

    add_h(doc, "G. Statistics, Latency, and ASK", level=2)
    add_p(
        doc,
        "Bootstrap resampling (n=1000) on limited-network / pilot_87 yields ASR 1.64% (95% CI "
        "[0.00, 5.45]), TSR 76.92% ([59.26, 92.00]), and F1 94.49% ([89.66, 98.41]). Decision-path "
        "latency on the measurement host is mean/median/p95/p99 = 0.75/0.68/1.27/1.89 ms. Pilot false "
        "positives under limited-network are 6/26 benign cases, all writes—an expected policy cost "
        "rather than an inference bug. Host for decision-path microbenchmarks: Windows 10, Python "
        "3.11.9; MCP serialization is excluded from those microbenchmarks and included in live "
        "end-to-end figures.",
    )
    add_p(
        doc,
        "The ASK study on approval-gated (Table XI) compares auto-DENY, auto-ALLOW, and author-as-user "
        "approval. Auto-DENY is safe (observed ASR 0%) but reduces TSR to 66.67%. Auto-ALLOW restores "
        "utility but reopens attacks (ASR 90.91%). Author-as-user approval restores TSR to 100% at "
        "observed ASR 0% with an approval rate of 23.1% and modest extra latency (50.4 ms), indicating "
        "that selective human confirmation can reclaim utility without broadly weakening the profile.",
    )
    add_table(
        doc,
        "TABLE XI. ASK / HUMAN-IN-THE-LOOP",
        ["Mode", "ASK #", "Approval", "ASR", "TSR", "F1", "Extra ms"],
        [
            ["auto_deny", "13", "0.0%", "0.00%", "66.67%", "88.00%", "0.0"],
            ["auto_allow", "13", "100%", "90.91%", "100%", "16.67%", "0.0"],
            ["author_ask", "13", "23.1%", "0.00%", "100%", "100%", "50.4"],
        ],
    )

    # VII Discussion
    add_h(doc, "VII. Discussion")
    add_h(doc, "A. Research Identity and Trade-off", level=2)
    add_p(
        doc,
        "The primary contribution is architectural: runtime, argument-aware, client-side "
        "least-privilege enforcement for MCP tool calls—not an ML detector. Across pilot, template "
        "stress, and Phase-C hold-outs, limited-network remains the strongest measured trade-off for "
        "interactive coding agents (reads + allow-listed HTTP). Read-only suits high-assurance review; "
        "no-code-exec suits edit+API workflows that can ban shells; approval-gated suits human-in-the-loop "
        "settings, as the ASK study suggests.",
    )
    add_h(doc, "B. Residual Risk, Unknown Tools, and Misconfiguration", level=2)
    add_p(
        doc,
        "Residual risk under fs.read-allowing profiles is mainly sensitive reads outside the current "
        "path list, and unconstrained net.http when domain allow-lists are omitted—both quantified by "
        "ablations. Mapping unknown tools to code.exec is fail-closed over-approximation: it explains "
        "the TSR collapse when name heuristics are removed and motivates a future first-class "
        "unknown policy category. Remapped names show that Layer 1 alone is brittle; argument "
        "inspection remains essential. The 20k corpus is stress-only; organic and live observed-ASR "
        "results carry the external-validity argument, within their sample-size limits.",
    )
    add_h(doc, "C. Deployment Implications", level=2)
    add_p(
        doc,
        "Because the proxy is configuration-only, organizations can introduce runtime least privilege "
        "without waiting for every MCP server vendor to ship manifests. Audit logs provide a practical "
        "incident-response trail for denied high-risk calls. The design is complementary to server "
        "sandboxes and gateway scanners: defense in depth remains advisable, but the client-side call "
        "boundary is where argument-aware least privilege can be enforced uniformly. Code, policies, "
        "datasets, and evaluation scripts are included in the submission artifact package "
        "(mcp_permission_layer/) for reproduction.",
    )
    add_h(doc, "D. Limitations and Future Work", level=2)
    add_p(
        doc,
        "Limitations: (i) taxonomy covers dominant studied behaviors, not all MCP surfaces (DB, "
        "browser, cloud IAM, email); (ii) unknown→code.exec is semantically coarse; (iii) template "
        "stress must not be over-interpreted; (iv) live (0/12) and adaptive (0/15) observed ASR=0% "
        "have one-sided 95% upper bounds ≤22.1% and ≤18.1%; (v) ASR is enforcement-boundary allow "
        "rate on curated calls, not end-to-end LLM agent compromise; (vi) multi-step full-chain "
        "success ignores partial-stage progress; (vii) session-taint is experimental only. Future "
        "work: first-class unknown category with explicit policy, larger live/adaptive corpora, "
        "coupling to live agent loops, richer argument semantics, production session tainting, and "
        "broader ASK user studies.",
    )

    # VIII Conclusion
    add_h(doc, "VIII. Conclusion")
    add_p(
        doc,
        "We presented a client-side runtime permission enforcement layer for MCP tool poisoning: a "
        "deterministic, argument-aware least-privilege gate—not an ML classifier—combining a "
        "behavior-mapped six-category model, fail-closed inference, declarative default-deny "
        "profiles, and a zero-modification stdio proxy (Algorithms 1–2).",
    )
    add_p(
        doc,
        "Under limited-network, enforcement-boundary ASR stays low while TSR remains high on pilot "
        "and Phase-C hold-outs, with sub-millisecond decision latency. Template stress supports "
        "scalability claims only; observed ASR 0% on small live/adaptive sets is reported with "
        "upper bounds. Baselines and ablations support structured runtime checks over name "
        "allowlists and definition-only scanning. Artifacts accompany the submission.",
    )

    # References
    add_h(doc, "References")
    refs = [
        "[1] K. Greshake, S. Abdelnabi, S. Mishra, C. Endres, T. Holz, and M. Fritz, “Not what you’ve signed up for: Compromising real-world LLM-integrated applications with indirect prompt injection,” in Proc. 16th ACM Workshop Artif. Intell. Secur. (AISec), 2023, pp. 79–90, doi: 10.1145/3605764.3623985.",
        "[2] Z. Wang et al., “MCPTox: A benchmark for tool poisoning attack on real-world MCP servers,” in Proc. AAAI Conf. Artif. Intell., vol. 40, no. 42, 2026, pp. 35811–35819, doi: 10.1609/aaai.v40i42.40895. Also arXiv:2508.14925.",
        "[3] R. Li, Z. Wang, Y. Yao, and X.-Y. Li, “MCP-ITP: An automated framework for implicit tool poisoning in MCP,” arXiv:2601.07395, 2026.",
        "[4] C. Bühler, M. Biagiola, L. Di Grazia, and G. Salvaneschi, “AgentBound: Securing execution boundaries of AI agents,” Proc. ACM Softw. Eng., vol. 3, no. FSE, Art. FSE096, 2026, doi: 10.1145/3808103. Also arXiv:2510.21236.",
        "[5] OWASP Foundation, “OWASP MCP Top 10,” OWASP Project. [Online]. Available: https://owasp.org/www-project-mcp-top-10/ (accessed Aug. 9, 2026).",
        "[6] Q. Zhan, Z. Liang, Z. Ying, and D. Kang, “InjecAgent: Benchmarking indirect prompt injections in tool-integrated large language model agents,” in Findings Assoc. Comput. Linguistics (ACL), 2024. Also arXiv:2403.02691.",
        "[7] E. Debenedetti et al., “AgentDojo: A dynamic environment to evaluate prompt injection attacks and defenses for LLM agents,” in Proc. NeurIPS Datasets and Benchmarks Track, 2024.",
        "[8] Microsoft Security, “Understanding and mitigating security risks in MCP implementations,” Microsoft Security Community Blog, 2025. [Online]. Available: https://techcommunity.microsoft.com/blog/microsoft-security-blog/understanding-and-mitigating-security-risks-in-mcp-implementations/4404667",
        "[9] R. Bhagwatkar et al., “Indirect prompt injections: Are firewalls all you need, or stronger benchmarks?” arXiv:2510.05244, 2025.",
        "[10] S. Yao et al., “ReAct: Synergizing reasoning and acting in language models,” in Proc. ICLR, 2023.",
        "[11] T. Schick et al., “Toolformer: Language models can teach themselves to use tools,” in Proc. NeurIPS, 2023.",
        "[12] Anthropic, “Introducing the Model Context Protocol,” Nov. 2024. [Online]. Available: https://www.anthropic.com/news/model-context-protocol. Spec.: https://modelcontextprotocol.io/",
        "[13] Android Developers, “Permissions on Android.” [Online]. Available: https://developer.android.com/guide/topics/permissions/overview",
        "[14] J. H. Saltzer and M. D. Schroeder, “The protection of information in computer systems,” Proc. IEEE, vol. 63, no. 9, pp. 1278–1308, 1975.",
        "[15] Microsoft for Developers, “Securing MCP: A control plane for agent tool execution,” 2026. [Online]. Available: https://developer.microsoft.com/blog/securing-mcp-a-control-plane-for-agent-tool-execution",
        "[16] E. Debenedetti et al., “Defeating prompt injections by design,” arXiv:2503.18813, 2025.",
        "[17] OWASP Foundation, “OWASP Top 10 for Large Language Model Applications,” version 2025. [Online]. Available: https://owasp.org/www-project-top-10-for-large-language-model-applications/",
        "[18] Anthropic, “Model Context Protocol (MCP) documentation,” Anthropic Docs. [Online]. Available: https://docs.anthropic.com/en/docs/agents-and-tools/mcp (accessed Aug. 9, 2026).",
        "[19] Cursor, “Model Context Protocol (MCP),” Cursor Docs. [Online]. Available: https://docs.cursor.com/context/model-context-protocol (accessed Aug. 9, 2026).",
        "[20] Cline, “MCP overview,” Cline Docs. [Online]. Available: https://docs.cline.bot/mcp/mcp-overview (accessed Aug. 9, 2026).",
        "[21] Continue, “How to set up Model Context Protocol (MCP) in Continue,” Continue Docs. [Online]. Available: https://docs.continue.dev/customize/deep-dives/mcp (accessed Aug. 9, 2026).",
        "[22] X. Hou, Y. Zhao, S. Wang, and H. Wang, “Model Context Protocol (MCP): Landscape, security threats, and future research directions,” arXiv:2503.23278, 2025.",
    ]
    for r in refs:
        add_p(doc, r, size=9, space_after=3, first_line=False)

    return doc


def estimate_pages(doc: Document) -> dict:
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    words = sum(len(t.split()) for t in paras)
    n_tables = len(doc.tables)
    n_figs = sum(1 for p in doc.paragraphs if p.text.strip().startswith("Fig."))
    n_algo_lines = sum(
        1
        for p in doc.paragraphs
        if p.runs
        and p.runs[0].font.name == "Courier New"
    )
    # Single-column 10/11pt with figures: conservative page model
    pages = words / 380.0 + n_tables * 0.35 + n_figs * 0.85 + n_algo_lines / 45.0 + 0.8
    return {
        "words": words,
        "tables": n_tables,
        "figs": n_figs,
        "algo_lines": n_algo_lines,
        "est_pages": round(pages, 1),
    }


def main():
    doc = build()
    stats = estimate_pages(doc)
    print(
        f"words={stats['words']} tables={stats['tables']} figs={stats['figs']} "
        f"est_pages~={stats['est_pages']}"
    )
    for out in OUTS:
        out.parent.mkdir(parents=True, exist_ok=True)
        candidates = [
            out,
            out.with_name(out.stem + "_SVG.docx"),
            out.with_name(out.stem + "_NEW.docx"),
            out.with_name(out.stem + "_v2.docx"),
        ]
        saved = False
        for candidate in candidates:
            try:
                doc.save(str(candidate))
                print("Wrote", candidate, "size", candidate.stat().st_size)
                saved = True
                break
            except PermissionError:
                continue
        if not saved:
            print("SKIP locked:", out)


if __name__ == "__main__":
    main()

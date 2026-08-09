"""Apply MDPI-style Figure/Table/Algorithm caption formatting to a Word paper."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

SRC = Path(r"c:\Users\as\Desktop\bodmas\IEEE_MCP_Paper_Style_Edited.docx")

ROMAN = {
    "I": 1,
    "II": 2,
    "III": 3,
    "IV": 4,
    "V": 5,
    "VI": 6,
    "VII": 7,
    "VIII": 8,
    "IX": 9,
    "X": 10,
    "XI": 11,
    "XII": 12,
}

ACRONYMS = [
    "MCP",
    "ASR",
    "TSR",
    "F1",
    "ASK",
    "YAML",
    "JSONL",
    "HTTP",
    "OWASP",
    "LLM",
    "API",
    "SSH",
    "IAM",
    "DB",
    "CI",
    "TP",
    "FP",
    "TN",
    "FN",
    "MAL",
    "DEC",
    "MS",
    "PILOT_87",
]


def roman_to_int(tok: str) -> int | None:
    tok = tok.upper()
    if tok.isdigit():
        return int(tok)
    return ROMAN.get(tok)


def sentence_case_title(title: str) -> str:
    t = title.strip().rstrip(".")
    letters = [c for c in t if c.isalpha()]
    upper_ratio = (sum(1 for c in letters if c.isupper()) / len(letters)) if letters else 0
    if upper_ratio >= 0.55:
        t = t.lower()
        for ac in ACRONYMS:
            t = re.sub(rf"\b{re.escape(ac.lower())}\b", ac, t, flags=re.I)
        t = re.sub(r"\bf1\b", "F1", t, flags=re.I)
        # Keep parenthetical abbreviations readable
        t = re.sub(r"\bms\b", "ms", t)
        t = t[:1].upper() + t[1:] if t else t
    if t and not t.endswith("."):
        t += "."
    return t


def style_run(run, *, bold: bool, size=9):
    run.bold = bold
    run.italic = False
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_caption(paragraph, kind: str, num: int, title: str):
    """MDPI style: bold 'Figure 1.' / 'Table 1.' then regular sentence-case title."""
    title = sentence_case_title(title)
    label = f"{kind} {num}."
    for r in paragraph.runs:
        r.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = label
        style_run(paragraph.runs[0], bold=True)
        for r in paragraph.runs[1:]:
            r.text = ""
        run2 = paragraph.add_run(f" {title}")
        style_run(run2, bold=False)
    else:
        r1 = paragraph.add_run(label)
        style_run(r1, bold=True)
        r2 = paragraph.add_run(f" {title}")
        style_run(r2, bold=False)


def replace_in_runs(paragraph, mapping: list[tuple[str, str]]):
    full = "".join(r.text for r in paragraph.runs) if paragraph.runs else paragraph.text
    new = full
    for a, b in mapping:
        new = new.replace(a, b)
    if new == full:
        return False
    if paragraph.runs:
        first = paragraph.runs[0]
        bold = first.bold
        size = first.font.size
        for r in paragraph.runs:
            r.text = ""
        first.text = new
        if bold is not None:
            first.bold = bold
        if size is not None:
            first.font.size = size
    else:
        paragraph.text = new
    return True


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    doc = Document(str(SRC))

    caption_table_re = re.compile(r"^(TABLE|Table)\s+([IVXLC]+|\d+)\.\s*(.*)$")
    caption_fig_re = re.compile(r"^(Fig\.|Figure)\s*(\d+)\.\s*(.*)$")
    caption_algo_re = re.compile(r"^(Algorithm)\s+(\d+)\s+(.+)$")

    changed: list[tuple[str, str, str]] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        m = caption_table_re.match(text)
        if m:
            n = roman_to_int(m.group(2))
            if n is None:
                continue
            before = text
            set_caption(p, "Table", n, m.group(3))
            changed.append(("table-caption", before, p.text))
            continue
        m = caption_fig_re.match(text)
        if m:
            before = text
            set_caption(p, "Figure", int(m.group(2)), m.group(3))
            changed.append(("fig-caption", before, p.text))
            continue
        m = caption_algo_re.match(text)
        if m and "formalises" not in text.lower() and "formalizes" not in text.lower():
            before = text
            set_caption(p, "Algorithm", int(m.group(2)), m.group(3))
            changed.append(("algo-caption", before, p.text))
            continue

    intext_map: list[tuple[str, str]] = []
    for rom, num in sorted(ROMAN.items(), key=lambda kv: -len(kv[0])):
        intext_map.append((f"Table {rom}", f"Table {num}"))
        intext_map.append((f"TABLE {rom}", f"Table {num}"))
    intext_map.append(("Fig. ", "Figure "))
    intext_map.append(("Figs. ", "Figures "))

    for p in doc.paragraphs:
        raw = p.text.strip()
        if not raw:
            continue
        if re.match(r"^(Table|Figure|Algorithm)\s+\d+\.\s", raw):
            continue
        before = p.text
        if replace_in_runs(p, intext_map):
            changed.append(("intext", before[:100], p.text[:100]))

    doc.save(str(SRC))
    print(f"Saved {SRC}")
    print(f"Changes: {len(changed)}")
    for kind, a, b in changed:
        print(f"[{kind}]")
        print("  FROM:", a[:170])
        print("  TO  :", b[:170])


if __name__ == "__main__":
    main()

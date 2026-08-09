"""Rewrite Algorithms 1–2 in formal MDPI (algorithm2e-like) academic style."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

SRC = Path(r"c:\Users\as\Desktop\bodmas\IEEE_MCP_Paper_Style_Edited.docx")

KEYWORDS = [
    "else if",
    "end for",
    "end if",
    "for each",
    "Require",
    "Ensure",
    "return",
    "else",
    "then",
    "true",
    "false",
    "do",
    "if",
]

ALGO1 = [
    "Algorithm 1. Permission category inference for an MCP tool call.",
    "Require: tool name n, argument map A",
    "Ensure: permission set P, high-risk flag h",
    "1:  P ← ∅; h ← false",
    "2:  // Layer 1: name-based heuristics",
    "3:  if n matches a filesystem-read pattern then P ← P ∪ {fs.read} end if",
    "4:  if n matches a filesystem-write pattern then P ← P ∪ {fs.write} end if",
    "5:  if n matches an HTTP-request pattern then P ← P ∪ {net.http} end if",
    "6:  if n matches a shell-execution pattern then P ← P ∪ {code.exec} end if",
    "7:  if n matches an environment-read pattern then P ← P ∪ {env.read} end if",
    "8:  if n matches a process-control pattern then P ← P ∪ {process} end if",
    "9:  // Layer 2: argument inspection",
    "10: for each (k, v) ∈ A do",
    "11:     if v denotes a filesystem path then",
    "12:         if writeContext(n) = true then P ← P ∪ {fs.write} else P ← P ∪ {fs.read} end if",
    "13:     end if",
    "14:     if v denotes a URL then P ← P ∪ {net.http} end if",
    "15:     if v contains shell metacharacters or shell commands then P ← P ∪ {code.exec} end if",
    "16:     if v matches environment-variable or secret keywords then P ← P ∪ {env.read} end if",
    "17:     if v matches high-risk host indicators then",
    "18:         P ← P ∪ {net.http, code.exec}; h ← true",
    "19:     end if",
    "20: end for",
    "21: // Layer 3: sensitive-path elevation",
    "22: for each path-valued argument v ∈ A do",
    "23:     if v matches a sensitive-path pattern (e.g., ~/.ssh/**, **/.env, **/*.pem) then",
    "24:         P ← P ∪ {env.read}; h ← true",
    "25:     end if",
    "26: end for",
    "27: if P = ∅ then P ← {code.exec}; h ← true end if  // fail-closed over-approximation",
    "28: return P, h",
]

ALGO2 = [
    "Algorithm 2. Policy evaluation and decision aggregation.",
    "Require: tool call τ, policy π",
    "Ensure: final decision d ∈ {ALLOW, DENY, ASK} and an audit record",
    "1:  (P, h) ← InferCategories(τ)  // Algorithm 1",
    "2:  D ← ∅",
    "3:  for each category c ∈ P do",
    "4:      r ← Lookup(π, c)",
    "5:      if r is undefined then r ← DefaultDeny end if",
    "6:      if r.allow = false then",
    "7:          D[c] ← DENY",
    "8:      else if the path or domain restrictions of r are violated then",
    "9:          D[c] ← DENY",
    "10:     else",
    "11:         D[c] ← r.mode  // ALLOW or ASK",
    "12:     end if",
    "13: end for",
    "14: // Aggregate with priority DENY ≻ ASK ≻ ALLOW",
    "15: if ∃ c ∈ P : D[c] = DENY then d ← DENY",
    "16: else if ∃ c ∈ P : D[c] = ASK then d ← ASK",
    "17: else d ← ALLOW end if",
    "18: if d = ASK and no ask-handler is configured then d ← DENY end if  // fail-closed",
    "19: Append an audit record for (τ, P, D, d)",
    "20: return d",
]


def style_run(run, *, bold=False, size=9, font="Courier New"):
    run.bold = bold
    run.italic = False
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)


def clear_runs(paragraph):
    for r in paragraph.runs:
        r.text = ""


def set_caption(paragraph, text: str):
    m = re.match(r"^(Algorithm\s+\d+\.)\s*(.*)$", text.strip())
    assert m
    clear_runs(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if paragraph.runs:
        paragraph.runs[0].text = m.group(1)
        style_run(paragraph.runs[0], bold=True, font="Times New Roman")
        for r in paragraph.runs[1:]:
            r.text = ""
        r2 = paragraph.add_run(" " + m.group(2))
        style_run(r2, bold=False, font="Times New Roman")
    else:
        r1 = paragraph.add_run(m.group(1))
        style_run(r1, bold=True, font="Times New Roman")
        r2 = paragraph.add_run(" " + m.group(2))
        style_run(r2, bold=False, font="Times New Roman")


def tokenize(text: str):
    pattern = re.compile(
        r"(?<![A-Za-z_])("
        + "|".join(re.escape(k) for k in KEYWORDS)
        + r")(?![A-Za-z_])",
        flags=re.IGNORECASE,
    )
    parts = []
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            parts.append((text[pos : m.start()], False))
        raw = m.group(1)
        canon = next(k for k in KEYWORDS if k.lower() == raw.lower())
        if canon in {"Require", "Ensure"}:
            out = canon
        else:
            out = canon.lower()
        parts.append((out, True))
        pos = m.end()
    if pos < len(text):
        parts.append((text[pos:], False))
    return parts


def set_code_line(paragraph, text: str):
    clear_runs(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    m = re.match(r"^(Require|Ensure):(.*)$", text)
    if m:
        chunks = [(m.group(1) + ":", True), (m.group(2), False)]
    else:
        chunks = tokenize(text)
    first = True
    for chunk, is_kw in chunks:
        if not chunk:
            continue
        if first and paragraph.runs:
            run = paragraph.runs[0]
            run.text = chunk
            first = False
        else:
            run = paragraph.add_run(chunk)
            first = False
        style_run(run, bold=is_kw, font="Courier New", size=9)


def insert_paragraph_after(paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def delete_paragraph(paragraph):
    p = paragraph._element
    parent = p.getparent()
    if parent is not None:
        parent.remove(p)


def find_algo_start(doc: Document, n: int) -> int:
    for i, p in enumerate(doc.paragraphs):
        if re.match(rf"^Algorithm\s+{n}\.", p.text.strip()):
            return i
    raise ValueError(f"Algorithm {n} caption not found")


def replace_algo_block(doc: Document, start_idx: int, lines: list[str], stop_before_algo: int | None):
    """Replace from caption at start_idx through algorithm body until prose/next algo."""
    # Determine end of current block
    end = start_idx + 1
    while end < len(doc.paragraphs):
        if stop_before_algo is not None and end >= stop_before_algo:
            break
        t = doc.paragraphs[end].text.strip()
        if not t:
            # peek next
            nxt = doc.paragraphs[end + 1].text.strip() if end + 1 < len(doc.paragraphs) else ""
            if nxt.startswith("Algorithm ") or (
                nxt
                and not nxt.startswith(("Require:", "Ensure:", "Input:", "Output:"))
                and not re.match(r"^\d+:", nxt)
                and not nxt.startswith("//")
                and not nxt.startswith("end ")
            ):
                break
            end += 1
            continue
        if t.startswith("Algorithm ") and re.match(r"^Algorithm\s+\d+\.", t):
            break
        if not (
            t.startswith(("Require:", "Ensure:", "Input:", "Output:"))
            or re.match(r"^\d+:", t)
            or t.startswith(("end if", "end for", "//"))
            or t.startswith("else")
        ):
            break
        end += 1

    caption_para = doc.paragraphs[start_idx]
    # Delete old body paragraphs (from end-1 down to start_idx+1)
    for i in range(end - 1, start_idx, -1):
        delete_paragraph(doc.paragraphs[i])

    # Caption
    set_caption(caption_para, lines[0])

    # Insert body lines after caption
    prev = caption_para
    for line in lines[1:]:
        para = insert_paragraph_after(prev)
        set_code_line(para, line)
        prev = para

    # trailing blank
    blank = insert_paragraph_after(prev)
    blank.paragraph_format.space_after = Pt(6)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    doc = Document(str(SRC))

    # Process Algorithm 2 first so indices for Algorithm 1 stay stable... 
    # Actually deleting changes indices; process from bottom.
    a2 = find_algo_start(doc, 2)
    replace_algo_block(doc, a2, ALGO2, stop_before_algo=None)

    a1 = find_algo_start(doc, 1)
    a2b = find_algo_start(doc, 2)
    replace_algo_block(doc, a1, ALGO1, stop_before_algo=a2b)

    # Soften intro sentence if it still says Input-style wording (optional leave)
    doc.save(str(SRC))
    print(f"Saved {SRC}")
    # verify
    d2 = Document(str(SRC))
    printing = False
    for i, p in enumerate(d2.paragraphs):
        t = p.text.strip()
        if t.startswith("Algorithm 1."):
            printing = True
        if printing:
            print(f"{i}|{t}")
        if t.startswith("Algorithm 2."):
            # continue through algo2
            pass
        if printing and t.startswith("The Python 3 prototype"):
            break
        if printing and t.startswith("IV.") or (printing and re.match(r"^[IVX]+\.", t)):
            # section after algorithms
            if i > 0 and d2.paragraphs[i - 1].text.strip().startswith("return"):
                pass


if __name__ == "__main__":
    main()

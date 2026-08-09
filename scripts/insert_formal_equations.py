"""Insert code/algorithm-consistent formal equations into the style-edited paper."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

SRC = Path(r"c:\Users\as\Desktop\bodmas\IEEE_MCP_Paper_Style_Edited.docx")


def style_run(run, *, bold=False, italic=False, size=10, font="Times New Roman"):
    run.bold = bold
    run.italic = italic
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)


def insert_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def set_text(
    paragraph: Paragraph,
    text: str,
    *,
    bold=False,
    italic=False,
    center=False,
    size=10,
    space_before=6,
    space_after=6,
    first_indent=True,
):
    for r in paragraph.runs:
        r.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = paragraph.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.line_spacing = 1.15
    pf.first_line_indent = Inches(0.25) if (first_indent and not center and not bold) else Pt(0)
    if paragraph.runs:
        paragraph.runs[0].text = text
        style_run(paragraph.runs[0], bold=bold, italic=italic, size=size)
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        r = paragraph.add_run(text)
        style_run(r, bold=bold, italic=italic, size=size)


# late import for Inches
from docx.shared import Inches  # noqa: E402


def add_block(after: Paragraph, items: list[tuple[str, dict]]) -> Paragraph:
    cur = after
    for text, kwargs in items:
        cur = insert_after(cur)
        set_text(cur, text, **kwargs)
    return cur


def find_para(doc: Document, prefix: str) -> tuple[int, Paragraph]:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(prefix):
            return i, p
    raise ValueError(prefix)


def find_return_line(doc: Document, start_idx: int, needle: str) -> Paragraph:
    for i in range(start_idx, len(doc.paragraphs)):
        if needle in doc.paragraphs[i].text:
            return doc.paragraphs[i]
    raise ValueError(needle)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    doc = Document(str(SRC))

    # Avoid double-insertion
    full = "\n".join(p.text for p in doc.paragraphs)
    if "Equation (1)." in full or "P = P_name ∪ P_args ∪ P_sens" in full:
        print("Equations already present; aborting to avoid duplication.")
        return

    # Update Algorithm intro to point to equations
    _, intro = find_para(doc, "Algorithm 1 formalises")
    intro_text = intro.text
    if "Equations (1)" not in intro_text:
        intro_text = intro_text.replace(
            "Both algorithms are implemented verbatim in the prototype and exercised by unit tests and the evaluation harness.",
            "Equations (1)–(9) state the same inference and decision rules in closed form, matching Algorithms 1–2 and the "
            "prototype functions infer_categories() and PermissionChecker.evaluate(); Equations (10)–(16) define the "
            "evaluation metrics computed from those decisions. Both algorithms are implemented "
            "verbatim in the prototype and exercised by unit tests and the evaluation harness.",
        )
        for r in intro.runs:
            r.text = ""
        intro.runs[0].text = intro_text
        style_run(intro.runs[0], size=10)

    # ---- After Algorithm 1 return ----
    i1, _ = find_para(doc, "Algorithm 1.")
    a1_ret = find_return_line(doc, i1, "return P, h")
    cur = add_block(
        a1_ret,
        [
            (
                "The inference procedure in Algorithm 1 is exactly the three-layer construction used in "
                "src/inference.py. Let n denote the tool name, A the argument map, P the inferred permission "
                "set, and h the high-risk flag. Layer outputs are combined by set union:",
                dict(space_before=8, space_after=4),
            ),
            (
                "P = P_name ∪ P_args ∪ P_sens  (1)",
                dict(center=True, first_indent=False, size=10, space_before=4, space_after=2),
            ),
            (
                "where P_name is produced by name heuristics (Algorithm 1, steps 3–8), P_args by argument "
                "inspection (steps 10–17), and P_sens by sensitive-path elevation (steps 19–21). Each layer "
                "only adds categories; it never removes a category inferred by an earlier layer. This matches "
                "the over-approximate safety design of the implementation.",
                dict(space_before=2, space_after=6),
            ),
            (
                "For a path-valued argument v ∈ A, write versus read is decided by the same writeContext(n) "
                "predicate used in code (is_write_operation):",
                dict(space_before=4, space_after=4),
            ),
            (
                "path(v) ∧ writeContext(n) = true  ⇒  P ← P ∪ {fs.write}  (2)",
                dict(center=True, first_indent=False, size=10, space_before=2, space_after=2),
            ),
            (
                "path(v) ∧ writeContext(n) = false ⇒  P ← P ∪ {fs.read}  (3)",
                dict(center=True, first_indent=False, size=10, space_before=2, space_after=4),
            ),
            (
                "URL-valued arguments contribute net.http; shell metacharacters or shell-command tokens "
                "contribute code.exec; environment-variable or secret keywords contribute env.read. High-risk "
                "host indicators simultaneously add {net.http, code.exec} and set h ← true. Sensitive-path "
                "matches (for example ~/.ssh/**, **/.env, **/*.pem) add env.read and also set h ← true.",
                dict(space_before=2, space_after=6),
            ),
            (
                "If no category is inferred, the fail-closed over-approximation used by the prototype is applied:",
                dict(space_before=4, space_after=4),
            ),
            (
                "if P = ∅ then P ← {code.exec} and h ← true  (4)",
                dict(center=True, first_indent=False, size=10, space_before=2, space_after=4),
            ),
            (
                "Equation (4) is an intentional semantic over-approximation of unknown tools as code.exec; "
                "it is not a claim that every unknown tool executes code, but a default-deny stance at the "
                "enforcement boundary.",
                dict(space_before=2, space_after=8),
            ),
        ],
    )

    # ---- After Algorithm 2 return ----
    # Re-find because indices shifted
    i2, _ = find_para(doc, "Algorithm 2.")
    a2_ret = find_return_line(doc, i2, "return d")
    # ensure we got Algorithm 2's return, not something else
    cur = add_block(
        a2_ret,
        [
            (
                "Algorithm 2 realises the policy lookup, restriction checks, and aggregation used in "
                "src/checker.py. After (P, h) ← InferCategories(τ), each category c ∈ P is mapped to a "
                "category-level decision D[c] ∈ {ALLOW, DENY, ASK}. Let r = Lookup(π, c), with missing rules "
                "replaced by DefaultDeny. Then:",
                dict(space_before=8, space_after=4),
            ),
            (
                "D[c] = DENY,  if r.allow = false or restrictions(r, τ) are violated  (5)",
                dict(center=True, first_indent=False, size=10, space_before=2, space_after=2),
            ),
            (
                "D[c] = r.mode, if r.allow = true and restrictions(r, τ) are satisfied  (6)",
                dict(center=True, first_indent=False, size=10, space_before=2, space_after=4),
            ),
            (
                "Here restrictions(r, τ) denotes the path allow/deny checks for fs.read, fs.write, and "
                "env.read, and the domain allow/deny checks for net.http, exactly as implemented by "
                "_evaluate_rule(). Categories without applicable restriction lists are treated as satisfied "
                "when r.allow = true, so D[c] = r.mode ∈ {ALLOW, ASK}.",
                dict(space_before=2, space_after=6),
            ),
            (
                "The final decision aggregates category decisions with fixed priority DENY ≻ ASK ≻ ALLOW, "
                "matching PermissionChecker._aggregate():",
                dict(space_before=4, space_after=4),
            ),
            (
                "d = DENY  if ∃ c ∈ P : D[c] = DENY  (7)",
                dict(center=True, first_indent=False, size=10, space_before=2, space_after=2),
            ),
            (
                "d = ASK   if no DENY and ∃ c ∈ P : D[c] = ASK  (8)",
                dict(center=True, first_indent=False, size=10, space_before=2, space_after=2),
            ),
            (
                "d = ALLOW otherwise  (9)",
                dict(center=True, first_indent=False, size=10, space_before=2, space_after=4),
            ),
            (
                "If d = ASK and no ask-handler is configured, the proxy fails closed with d ← DENY. The audit "
                "record stores (τ, P, h, D, d) and decision-path latency, corresponding to DecisionRecord in "
                "the prototype. Equations (5)–(9) are therefore not an alternative model; they are the closed-form "
                "statement of Algorithm 2 and of the deployed checker.",
                dict(space_before=2, space_after=8),
            ),
        ],
    )

    # ---- Expand metrics section ----
    _, metrics = find_para(doc, "The evaluation reports enforcement-boundary metrics")
    # Replace metrics paragraph with shorter lead-in, then add equations after it
    lead = (
        "The evaluation reports enforcement-boundary metrics rather than full end-to-end agent-compromise "
        "metrics. All rates below are computed on the mediator decision d ∈ {ALLOW, DENY, ASK}, with ASK "
        "mapped to DENY in automated non-ASK studies, exactly as in the evaluation harness. Let M be the "
        "malicious case set and B the benign case set under a given study split."
    )
    for r in metrics.runs:
        r.text = ""
    metrics.runs[0].text = lead
    style_run(metrics.runs[0], size=10)

    cur = add_block(
        metrics,
        [
            (
                "Attack Success Rate (ASR) is the fraction of malicious tool calls that the layer allows to "
                "the upstream boundary:",
                dict(space_before=6, space_after=4),
            ),
            (
                "ASR = |{ τ ∈ M : d(τ) = ALLOW }| / |M|  (10)",
                dict(center=True, first_indent=False, size=10, space_before=2, space_after=4),
            ),
            (
                "Block Rate on the malicious subset is the complement of ASR and equals the true-positive "
                "rate of attack blocking:",
                dict(space_before=2, space_after=4),
            ),
            (
                "Block = 1 − ASR = |{ τ ∈ M : d(τ) = DENY }| / |M|  (11)",
                dict(center=True, first_indent=False, size=10, space_before=2, space_after=4),
            ),
            (
                "Task Success Rate (TSR) is the fraction of benign tool calls that remain allowed:",
                dict(space_before=2, space_after=4),
            ),
            (
                "TSR = |{ τ ∈ B : d(τ) = ALLOW }| / |B|  (12)",
                dict(center=True, first_indent=False, size=10, space_before=2, space_after=4),
            ),
            (
                "Attack-Blocking F1 treats a blocked attack as the positive class. With",
                dict(space_before=2, space_after=4),
            ),
            (
                "TP = |{ τ ∈ M : d(τ) = DENY }|,   FP = |{ τ ∈ B : d(τ) = DENY }|  (13)",
                dict(center=True, first_indent=False, size=10, space_before=2, space_after=2),
            ),
            (
                "FN = |{ τ ∈ M : d(τ) = ALLOW }|,   TN = |{ τ ∈ B : d(τ) = ALLOW }|  (14)",
                dict(center=True, first_indent=False, size=10, space_before=2, space_after=2),
            ),
            (
                "F1 = 2·TP / (2·TP + FP + FN)  (15)",
                dict(center=True, first_indent=False, size=10, space_before=2, space_after=4),
            ),
            (
                "These definitions are identical to the counts returned by the evaluation scripts "
                "(confusion_counts / derive_rates). In particular, FN corresponds to residual ASR cases, "
                "and FP corresponds to blocked benign utility loss. Latency is reported as mean, median, "
                "p95, and p99 on the decision path; live MCP additionally reports end-to-end latency including "
                "upstream execution. Bootstrap 95% confidence intervals (n=1000) are reported for "
                "limited-network / pilot_87. For observed ASR = 0 on a hold-out of size n, the one-sided "
                "95% upper bound used in the text is",
                dict(space_before=2, space_after=4),
            ),
            (
                "ASR_UB = 1 − 0.05^(1/n)  (16)",
                dict(center=True, first_indent=False, size=10, space_before=2, space_after=4),
            ),
            (
                "Baselines are B0 none, B1 static allowlist, B2 definition scan, B3 gateway filter, and B4 "
                "the full layer. All automated experiments use the same harness entry points so that policies, "
                "ablation flags, and dataset loaders remain directly comparable.",
                dict(space_before=2, space_after=8),
            ),
        ],
    )

    # Trim duplicate follow-on paragraph if it largely repeats baselines/latency
    # Find the old next paragraph after metrics block - it may still be the long original continuation
    # After our insertion, the paragraph that used to follow metrics may still have old duplicate content.
    # Locate "All automated experiments use the same harness"
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("All automated experiments use the same harness entry points"):
            # Keep figure intro only if present later; shorten this duplicate
            # If our new metrics already covered harness sentence, replace with live-MCP-only sentence
            new = (
                "Live MCP experiments exercise the real stdio proxy path rather than a mocked checker, "
                "ensuring that serialization, forwarding, and denial short-circuiting are included in the "
                "end-to-end measurements."
            )
            if "Live MCP experiments" in t and "Baselines are B0" not in t:
                # already the short live paragraph - leave
                pass
            else:
                for r in p.runs:
                    r.text = ""
                if p.runs:
                    p.runs[0].text = new
                    style_run(p.runs[0], size=10)
                else:
                    r = p.add_run(new)
                    style_run(r, size=10)
            break

    doc.save(str(SRC))
    print("Saved", SRC)

    # verify
    d2 = Document(str(SRC))
    text = "\n".join(p.text for p in d2.paragraphs)
    for eq in range(1, 17):
        ok = f"({eq})" in text or f"({eq})." in text
        # our format uses  (n)
        ok = f"({eq})" in text
        print(f"Eq ({eq}):", "OK" if ok else "MISSING")
    for needle in [
        "P = P_name ∪ P_args ∪ P_sens",
        "infer_categories()",
        "_aggregate()",
        "ASR = ",
        "F1 = 2·TP",
        "ASR_UB",
    ]:
        print(("OK" if needle in text else "MISSING"), needle)


if __name__ == "__main__":
    main()
